"""Safe preview + commit pipeline for team imports."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Iterable

from openpyxl import load_workbook

from .admin_repository import _has_tournament_access
from .repository import all_rows, connect

_HEADER_ALIASES = {
    "name": {"name", "team", "team name", "lag", "lagnamn", "lag namn"},
    "age_class": {"age class", "age_class", "ageclass", "åldersklass", "aldersklass", "klass"},
    "group": {"group", "group name", "grupp", "gruppnamn", "grupp namn"},
    "primary_color": {"primary color", "primary_color", "primärfärg", "primarfarg", "färg", "farg"},
    "secondary_color": {"secondary color", "secondary_color", "sekundärfärg", "sekundarfarg"},
}


def _normalize_header(value) -> str:
    return " ".join(str(value or "").strip().lower().replace("-", " ").replace("_", " ").split())


def _canonical_header(value):
    normalized = _normalize_header(value)
    for field, aliases in _HEADER_ALIASES.items():
        if normalized in aliases:
            return field
    return None


def _rows_from_csv(content: bytes) -> list[list]:
    text = content.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if ";" in sample else ","
    return [list(row) for row in csv.reader(StringIO(text), dialect)]


def _rows_from_xlsx(content: bytes) -> list[list]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def _parse(content: bytes, filename: str) -> list[dict]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        raw = _rows_from_csv(content)
    elif lower.endswith(".xlsx"):
        raw = _rows_from_xlsx(content)
    else:
        raise ValueError("Endast CSV och XLSX stöds")
    raw = [row for row in raw if any(str(value or "").strip() for value in row)]
    if not raw:
        raise ValueError("Filen innehåller inga rader")
    headers = [_canonical_header(value) for value in raw[0]]
    if "name" not in headers:
        raise ValueError("Ingen kolumn för lagnamn hittades. Använd till exempel 'Lag' eller 'Lagnamn'.")
    mapped = []
    for line_no, row in enumerate(raw[1:], start=2):
        item = {"row": line_no}
        for index, field in enumerate(headers):
            if field and index < len(row):
                value = row[index]
                item[field] = str(value).strip() if value is not None else ""
        mapped.append(item)
    return mapped


def preview_team_import(account_id: int, tournament_id: int, content: bytes, filename: str):
    if not _has_tournament_access(account_id, tournament_id):
        return None
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Filen är större än 5 MB")
    parsed = _parse(content, filename)
    existing = {str(row["name"]).strip().casefold() for row in all_rows("SELECT name FROM teams WHERE tournament_id=?", (int(tournament_id),))}
    groups = all_rows("SELECT id,name FROM groups WHERE tournament_id=?", (int(tournament_id),))
    group_by_name = {str(row["name"]).strip().casefold(): row for row in groups}
    seen = set()
    rows = []
    for raw in parsed:
        name = str(raw.get("name") or "").strip()
        group_name = str(raw.get("group") or "").strip()
        errors = []
        warnings = []
        if not name:
            errors.append("Lagnamn saknas")
        folded = name.casefold()
        if name and folded in seen:
            errors.append("Dubblett i filen")
        if name:
            seen.add(folded)
        if name and folded in existing:
            errors.append("Laget finns redan i cupen")
        group_id = None
        if group_name:
            group = group_by_name.get(group_name.casefold())
            if not group:
                errors.append(f"Gruppen '{group_name}' finns inte i cupen")
            else:
                group_id = int(group["id"])
        primary = str(raw.get("primary_color") or "").strip() or "#111827"
        secondary = str(raw.get("secondary_color") or "").strip() or "#FFFFFF"
        for label, color in (("Primärfärg", primary), ("Sekundärfärg", secondary)):
            if len(color) != 7 or not color.startswith("#") or any(c not in "0123456789abcdefABCDEF" for c in color[1:]):
                errors.append(f"{label} måste anges som #RRGGBB")
        rows.append({
            "row": raw["row"], "name": name, "age_class": str(raw.get("age_class") or "").strip() or None,
            "group": group_name or None, "group_id": group_id, "primary_color": primary,
            "secondary_color": secondary, "errors": errors, "warnings": warnings, "valid": not errors,
        })
    return {"filename": filename, "rows": rows, "valid_count": sum(1 for r in rows if r["valid"]), "error_count": sum(1 for r in rows if not r["valid"]), "groups": groups}


def commit_team_import(account_id: int, tournament_id: int, rows: Iterable[dict]):
    if not _has_tournament_access(account_id, tournament_id):
        return None
    rows = list(rows)
    if not rows:
        raise ValueError("Det finns inga lag att importera")
    if len(rows) > 1000:
        raise ValueError("Högst 1000 lag kan importeras åt gången")
    existing = {str(row["name"]).strip().casefold() for row in all_rows("SELECT name FROM teams WHERE tournament_id=?", (int(tournament_id),))}
    group_ids = {int(row["id"]) for row in all_rows("SELECT id FROM groups WHERE tournament_id=?", (int(tournament_id),))}
    seen = set()
    clean = []
    for index, row in enumerate(rows, start=1):
        name = str(row.get("name") or "").strip()
        if not name:
            raise ValueError(f"Rad {index}: lagnamn saknas")
        key = name.casefold()
        if key in existing or key in seen:
            raise ValueError(f"Rad {index}: laget '{name}' finns redan")
        seen.add(key)
        group_id = row.get("group_id")
        group_id = int(group_id) if group_id is not None else None
        if group_id is not None and group_id not in group_ids:
            raise ValueError(f"Rad {index}: gruppen finns inte längre")
        primary = str(row.get("primary_color") or "#111827").strip()
        secondary = str(row.get("secondary_color") or "#FFFFFF").strip()
        for color in (primary, secondary):
            if len(color) != 7 or not color.startswith("#") or any(c not in "0123456789abcdefABCDEF" for c in color[1:]):
                raise ValueError(f"Rad {index}: ogiltig lagfärg")
        clean.append((int(tournament_id), name, group_id, str(row.get("age_class") or "").strip() or None, primary, secondary))
    with connect() as con:
        try:
            for values in clean:
                con.execute("INSERT INTO teams(tournament_id,name,group_id,age_class,primary_color,secondary_color) VALUES(?,?,?,?,?,?)", values)
            commit = getattr(con, "commit", None)
            if callable(commit): commit()
        except Exception:
            rollback = getattr(con, "rollback", None)
            if callable(rollback): rollback()
            raise
    return {"imported": len(clean)}
